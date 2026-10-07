import openpyxl, sys, re, collections
crm, fsr, out = sys.argv[1], sys.argv[2], sys.argv[3]
TEAM = {'FS West':'FS WEST','FS East':'FS EAST','FS Central':'FS Central','FS South':'FS South'}
f = openpyxl.load_workbook(fsr, read_only=True, data_only=True)
members = []
for r in f['담당자 명단'].iter_rows(min_row=6, values_only=True):
    if r[1] and r[2] in TEAM: members.append((TEAM[r[2]], r[1].strip()))
accts = []
for r in f['담당기관 마스터'].iter_rows(min_row=5, values_only=True):
    if r[5] and r[3] in TEAM: accts.append((str(r[5]).strip(), TEAM[r[3]]))
RULES = [  # (regex on model, group)
 (r'^HISCL', 'HISCL'), (r'^HLC-?723', 'A1c'), (r'^RF-500', 'RF-500'),
 (r'^(CA|CS|CN)-', 'COA'),
 (r'^(UF|UC|UN|UD)-|^U-WAM', 'Urine'),
 (r'^(XN|XR|XS|XT|XE|XP|XQ|KX|K-|SP|DI|CF|BT|SA|RU|DC|CV|TS)-?', 'CBC'),
]
w = openpyxl.load_workbook(crm, read_only=True)
models = collections.Counter()
for r in w['2. 모든 약속_활동'].iter_rows(min_row=2, values_only=True):
    s = r[8]
    if s:
        p = s.split('/')
        if len(p) >= 3 and p[1]: models[p[1]] += 1
ig = []
for m, n in models.most_common():
    for rx, g in RULES:
        if re.match(rx, m): ig.append((m, g, n)); break
wb = openpyxl.Workbook()
ws = wb.active; ws.title = 'Members'
for t, n in members: ws.append([t, n])
ws = wb.create_sheet('ITEM Group')
for m, g, n in ig: ws.append([None, m, None, g, n])
ws = wb.create_sheet('거래처 List')
ws.append(['거래처 List (FSR 담당기관 관리표 기준)']); ws.append(['거래처', None, None, None, None, None, None, '담당팀'])
for a, t in accts: ws.append([a, None, None, None, None, None, None, t])
wb.save(out)
print('members', len(members), collections.Counter(t for t,_ in members))
print('item groups', len(ig), collections.Counter(g for _,g,_ in ig))
print('accounts', len(accts))
